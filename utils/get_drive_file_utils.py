import json
import math
import os
import re
import time
from datetime import date, datetime, time as datetime_time
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

from google.oauth2 import service_account
from google.auth.transport.requests import AuthorizedSession
from openpyxl import load_workbook

from database import db
from models import DriveFileDownload
from logging_config import logger


# ============================================================
# CONFIGURACION
# ============================================================

DEFAULT_DRIVE_URL = os.getenv(
    "GOOGLE_DRIVE_FILE_URL",
    (
        "https://docs.google.com/spreadsheets/d/"
        "1uzK9UQOIW5JQD-4VR-paWrKTvpZM2ao_IGEGK0QSg0o/"
        "edit?gid=1374924463#gid=1374924463"
    )
)

DEFAULT_FILE_NAME = os.getenv(
    "GOOGLE_DRIVE_FILE_NAME",
    "Rating_YPF.xlsx"
)

DEFAULT_HISTORY_URL = os.getenv(
    "GOOGLE_DRIVE_HISTORY_URL",
    ""
)

DEFAULT_HISTORY_FILE_NAME = os.getenv(
    "GOOGLE_DRIVE_HISTORY_FILE_NAME",
    "Rating_YPF_Historico.xlsx"
)

TARGET_SHEET_NAME = os.getenv(
    "GOOGLE_DRIVE_TARGET_SHEET_NAME",
    "Comentarios x orden"
)

EXPECTED_HEADERS = (
    "order_id",
    "survey_date",
    "partner_id",
    "partner_name",
    "product_name",
    "dish_rating",
    "dish_pill",
    "partner_comment"
)

GOOGLE_SERVICE_ACCOUNT_JSON_ENV = "GOOGLE_SERVICE_ACCOUNT_JSON"

DRIVE_READONLY_SCOPE = "https://www.googleapis.com/auth/drive.readonly"
SHEETS_SCOPE = "https://www.googleapis.com/auth/spreadsheets"

GOOGLE_SCOPES = [
    DRIVE_READONLY_SCOPE,
    SHEETS_SCOPE
]

XLSX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


# ============================================================
# GOOGLE AUTH / IDS
# ============================================================

def extract_google_file_id(url):
    """
    Extrae el ID de una URL de Google Drive / Google Sheets.
    """

    match = re.search(
        r"/d/([a-zA-Z0-9_-]+)",
        url or ""
    )

    if not match:
        raise ValueError(
            "No se pudo obtener el ID del archivo "
            "desde la URL de Google Drive"
        )

    return match.group(1)


def build_google_credentials():
    """
    Lee GOOGLE_SERVICE_ACCOUNT_JSON.

    La misma Service Account:
    - lee el Google Sheet origen;
    - escribe el Google Sheet historico.
    """

    raw_credentials = os.getenv(
        GOOGLE_SERVICE_ACCOUNT_JSON_ENV
    )

    if not raw_credentials:
        raise RuntimeError(
            "Falta la variable de entorno "
            "GOOGLE_SERVICE_ACCOUNT_JSON"
        )

    try:
        service_account_info = json.loads(raw_credentials)

    except json.JSONDecodeError as e:
        raise RuntimeError(
            "GOOGLE_SERVICE_ACCOUNT_JSON no contiene "
            "un JSON valido. "
            f"Detalle: {str(e)}"
        ) from e

    if service_account_info.get("type") != "service_account":
        raise RuntimeError(
            "GOOGLE_SERVICE_ACCOUNT_JSON no parece "
            "ser una credencial de Service Account"
        )

    required_fields = (
        "project_id",
        "private_key",
        "client_email",
        "token_uri"
    )

    missing_fields = [
        field
        for field in required_fields
        if not service_account_info.get(field)
    ]

    if missing_fields:
        raise RuntimeError(
            "Faltan campos obligatorios en "
            "GOOGLE_SERVICE_ACCOUNT_JSON: "
            + ", ".join(missing_fields)
        )

    try:
        return (
            service_account.Credentials
            .from_service_account_info(
                service_account_info,
                scopes=GOOGLE_SCOPES
            )
        )

    except Exception as e:
        raise RuntimeError(
            "No se pudieron construir las credenciales "
            f"de Google: {str(e)}"
        ) from e


# ============================================================
# DRIVE API - DESCARGA XLSX CON files.download
# ============================================================

def build_drive_download_url(drive_file_id):
    return (
        "https://www.googleapis.com/drive/v3/files/"
        f"{drive_file_id}/download"
    )


def build_operation_url(operation_name):
    operation_id = operation_name.split("/")[-1]

    return (
        "https://www.googleapis.com/drive/v3/operations/"
        f"{operation_id}"
    )


def download_google_file_as_xlsx_by_id(
    drive_file_id,
    session
):
    """
    Descarga un Google Sheet como XLSX usando files.download
    y espera la Long Running Operation.
    """

    response = session.post(
        build_drive_download_url(drive_file_id),
        params={
            "mimeType": XLSX_MIME_TYPE
        },
        headers={
            "Accept": "application/json",
            "Content-Length": "0"
        },
        timeout=(15, 60)
    )

    if response.status_code not in (200, 201):
        raise RuntimeError(
            "Google Drive API rechazo el inicio de la descarga. "
            f"HTTP {response.status_code}. "
            f"Detalle: {response.text[:2000]}"
        )

    try:
        operation_data = response.json()

    except Exception as e:
        raise RuntimeError(
            "Google Drive no devolvio una operacion JSON valida"
        ) from e

    operation_name = operation_data.get("name")

    if not operation_name:
        raise RuntimeError(
            "Google Drive no devolvio el nombre de la operacion. "
            f"Respuesta: {operation_data}"
        )

    operation_url = build_operation_url(operation_name)

    max_attempts = 150
    wait_seconds = 2
    final_operation = None

    for _attempt in range(1, max_attempts + 1):

        if operation_data.get("done"):
            final_operation = operation_data
            break

        time.sleep(wait_seconds)

        operation_response = session.get(
            operation_url,
            timeout=(15, 60)
        )

        if operation_response.status_code != 200:
            raise RuntimeError(
                "Google Drive API rechazo la consulta "
                "del estado de la operacion. "
                f"HTTP {operation_response.status_code}. "
                f"Detalle: {operation_response.text[:2000]}"
            )

        try:
            operation_data = operation_response.json()

        except Exception as e:
            raise RuntimeError(
                "Google Drive devolvio una respuesta invalida "
                "al consultar la operacion"
            ) from e

    if operation_data.get("done"):
        final_operation = operation_data

    if not final_operation:
        raise TimeoutError(
            "Google Drive tardo demasiado en preparar "
            "la descarga del archivo"
        )

    if final_operation.get("error"):
        raise RuntimeError(
            "Google Drive fallo al preparar la descarga. "
            f"Detalle: {final_operation['error']}"
        )

    download_uri = final_operation.get(
        "response",
        {}
    ).get("downloadUri")

    if not download_uri:
        raise RuntimeError(
            "Google Drive termino la operacion pero "
            "no devolvio downloadUri. "
            f"Respuesta: {final_operation}"
        )

    file_response = session.get(
        download_uri,
        stream=True,
        allow_redirects=True,
        timeout=(15, 300)
    )

    if file_response.status_code != 200:
        raise RuntimeError(
            "No se pudo descargar el XLSX generado. "
            f"HTTP {file_response.status_code}. "
            f"Detalle: {file_response.text[:2000]}"
        )

    content_type = file_response.headers.get(
        "Content-Type",
        ""
    ).lower()

    if "text/html" in content_type:
        raise RuntimeError(
            "Google devolvio HTML en lugar del XLSX. "
            "Revisar permisos del archivo y la Service Account."
        )

    buffer = BytesIO()

    for chunk in file_response.iter_content(
        chunk_size=1024 * 1024
    ):
        if chunk:
            buffer.write(chunk)

    binary_data = buffer.getvalue()

    if not binary_data:
        raise RuntimeError(
            "Google Drive devolvio un archivo vacio"
        )

    if not binary_data.startswith(b"PK"):
        raise RuntimeError(
            "La respuesta de Google no parece ser un XLSX valido"
        )

    return binary_data


def download_google_sheet_as_xlsx(drive_url):
    """
    Descarga cualquier Google Sheet compartido con la
    Service Account como XLSX.
    """

    drive_file_id = extract_google_file_id(drive_url)

    credentials = build_google_credentials()
    session = AuthorizedSession(credentials)

    logger.info(
        "Descargando Google Sheet. "
        "drive_file_id=%s service_account=%s",
        drive_file_id,
        credentials.service_account_email
    )

    binary_data = download_google_file_as_xlsx_by_id(
        drive_file_id,
        session
    )

    logger.info(
        "Google Sheet descargado correctamente. "
        "drive_file_id=%s size=%s bytes",
        drive_file_id,
        len(binary_data)
    )

    return {
        "drive_file_id": drive_file_id,
        "data": binary_data,
        "size_bytes": len(binary_data)
    }


# ============================================================
# XLSX - TRABAJAR SOLO CON "Comentarios x orden"
# ============================================================

def _find_sheet_name(workbook, wanted_name):
    if wanted_name in workbook.sheetnames:
        return wanted_name

    wanted_normalized = wanted_name.strip().casefold()

    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().casefold() == wanted_normalized:
            return sheet_name

    raise RuntimeError(
        f"No existe la hoja '{wanted_name}' en el archivo. "
        f"Hojas disponibles: {', '.join(workbook.sheetnames)}"
    )


def _normalize_header_value(value):
    if value is None:
        return ""

    return str(value).strip().casefold()


def _is_expected_header_row(row):
    """
    Detecta la fila real de encabezados de "Comentarios x orden".

    El archivo origen puede traer antes una fila de metadata como:
    "Custom query: ... Last updated: ..."

    Esa fila NO forma parte de los datos y debe descartarse.
    """

    normalized = [
        _normalize_header_value(value)
        for value in list(row)[:len(EXPECTED_HEADERS)]
    ]

    expected = [
        header.casefold()
        for header in EXPECTED_HEADERS
    ]

    return normalized == expected


def _find_header_row_number(worksheet):
    """
    Devuelve el numero de fila (1-based) donde estan los encabezados.

    Busca solamente en las primeras 25 filas para tolerar metadata
    o filas vacias antes del encabezado real.

    Importante:
    en workbooks abiertos con read_only=True, openpyxl puede devolver
    worksheet.max_row = None y worksheet.max_column = None.
    Por eso NO usamos max_row/max_column para determinar el rango.
    """

    for row_number, raw_row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=25,
            min_col=1,
            max_col=len(EXPECTED_HEADERS),
            values_only=True
        ),
        start=1
    ):
        values = list(raw_row)

        if _is_expected_header_row(values):
            return row_number

    raise RuntimeError(
        f"No se encontro el encabezado esperado en la hoja "
        f"'{worksheet.title}'. Se esperaban las columnas: "
        + ", ".join(EXPECTED_HEADERS)
    )


def keep_only_target_sheet(
    xlsx_binary,
    sheet_name=TARGET_SHEET_NAME
):
    """
    Devuelve un XLSX que:

    - conserva solamente la hoja objetivo;
    - elimina metadata/filas que aparezcan antes del encabezado real;
    - deja los encabezados en la fila 1.
    """

    try:
        workbook = load_workbook(
            BytesIO(xlsx_binary)
        )

    except Exception as e:
        raise RuntimeError(
            "No se pudo abrir el XLSX descargado "
            f"para filtrar la hoja. Detalle: {str(e)}"
        ) from e

    try:
        real_sheet_name = _find_sheet_name(
            workbook,
            sheet_name
        )

        for worksheet in list(workbook.worksheets):
            if worksheet.title != real_sheet_name:
                workbook.remove(worksheet)

        worksheet = workbook[real_sheet_name]

        header_row_number = _find_header_row_number(
            worksheet
        )

        if header_row_number > 1:
            rows_to_delete = header_row_number - 1

            logger.info(
                "Eliminando metadata previa al encabezado. "
                "sheet=%s rows=%s",
                real_sheet_name,
                rows_to_delete
            )

            worksheet.delete_rows(
                1,
                rows_to_delete
            )

        output_buffer = BytesIO()
        workbook.save(output_buffer)

    finally:
        workbook.close()

    filtered_binary = output_buffer.getvalue()

    if not filtered_binary.startswith(b"PK"):
        raise RuntimeError(
            "El XLSX filtrado no parece ser valido"
        )

    return filtered_binary


def _google_safe_value(value):
    """
    Convierte valores de openpyxl a tipos que JSON pueda enviar
    a Google Sheets.
    """

    if value is None:
        return ""

    if isinstance(
        value,
        (datetime, date, datetime_time)
    ):
        return value.isoformat()

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return str(value)

    if isinstance(
        value,
        (str, int, float, bool)
    ):
        return value

    return str(value)


def _trim_row(row):
    """
    Quita celdas vacias solamente del final de la fila.
    """

    result = list(row)

    while result and (
        result[-1] is None
        or result[-1] == ""
    ):
        result.pop()

    return result


def extract_target_sheet_rows(
    xlsx_binary,
    sheet_name=TARGET_SHEET_NAME
):
    """
    Lee la hoja objetivo ignorando cualquier metadata previa.

    La primera fila devuelta siempre es el encabezado real:
    order_id, survey_date, ..., partner_comment.
    """

    workbook = load_workbook(
        BytesIO(xlsx_binary),
        read_only=True,
        data_only=True
    )

    try:
        real_sheet_name = _find_sheet_name(
            workbook,
            sheet_name
        )

        worksheet = workbook[real_sheet_name]

        header_row_number = _find_header_row_number(
            worksheet
        )

        rows = []

        for raw_row in worksheet.iter_rows(
            min_row=header_row_number,
            values_only=True
        ):
            row = _trim_row(raw_row)

            if not row:
                continue

            if all(
                value is None or value == ""
                for value in row
            ):
                continue

            rows.append([
                _google_safe_value(value)
                for value in row
            ])

    finally:
        workbook.close()

    if not rows:
        raise RuntimeError(
            f"La hoja '{sheet_name}' no contiene datos"
        )

    if not _is_expected_header_row(rows[0]):
        raise RuntimeError(
            f"La hoja '{sheet_name}' no tiene el encabezado esperado"
        )

    return rows


# ============================================================
# SHEETS API - HISTORICO
# ============================================================

def _quoted_sheet_name(sheet_name):
    escaped = sheet_name.replace("'", "''")
    return f"'{escaped}'"


def _sheet_values_url(
    spreadsheet_id,
    sheet_name
):
    a1_range = _quoted_sheet_name(sheet_name)
    encoded_range = quote(a1_range, safe="")

    return (
        "https://sheets.googleapis.com/v4/spreadsheets/"
        f"{spreadsheet_id}/values/{encoded_range}"
    )


def _sheet_append_url(
    spreadsheet_id,
    sheet_name
):
    return (
        _sheet_values_url(
            spreadsheet_id,
            sheet_name
        )
        + ":append"
    )


def ensure_history_sheet_exists(
    session,
    spreadsheet_id,
    sheet_name=TARGET_SHEET_NAME
):
    """
    Si el historico esta nuevo y tiene una sola pestaña,
    la renombra. Si hay varias y falta la objetivo, la crea.
    """

    metadata_url = (
        "https://sheets.googleapis.com/v4/spreadsheets/"
        f"{spreadsheet_id}"
    )

    response = session.get(
        metadata_url,
        params={
            "fields": "sheets.properties(sheetId,title)"
        },
        timeout=(15, 60)
    )

    if response.status_code != 200:
        raise RuntimeError(
            "No se pudo leer el Google Sheet historico. "
            f"HTTP {response.status_code}. "
            f"Detalle: {response.text[:2000]}"
        )

    properties = [
        item.get("properties", {})
        for item in response.json().get("sheets", [])
    ]

    if any(
        props.get("title") == sheet_name
        for props in properties
    ):
        return

    batch_update_url = (
        "https://sheets.googleapis.com/v4/spreadsheets/"
        f"{spreadsheet_id}:batchUpdate"
    )

    if len(properties) == 1:
        payload = {
            "requests": [
                {
                    "updateSheetProperties": {
                        "properties": {
                            "sheetId": properties[0]["sheetId"],
                            "title": sheet_name
                        },
                        "fields": "title"
                    }
                }
            ]
        }

    else:
        payload = {
            "requests": [
                {
                    "addSheet": {
                        "properties": {
                            "title": sheet_name
                        }
                    }
                }
            ]
        }

    response = session.post(
        batch_update_url,
        json=payload,
        timeout=(15, 60)
    )

    if response.status_code != 200:
        raise RuntimeError(
            "No se pudo crear/renombrar la hoja "
            f"'{sheet_name}' en el historico. "
            f"HTTP {response.status_code}. "
            f"Detalle: {response.text[:2000]}"
        )


def get_history_rows(
    session,
    spreadsheet_id,
    sheet_name=TARGET_SHEET_NAME
):
    response = session.get(
        _sheet_values_url(
            spreadsheet_id,
            sheet_name
        ),
        params={
            "majorDimension": "ROWS",
            "valueRenderOption": "UNFORMATTED_VALUE",
            "dateTimeRenderOption": "FORMATTED_STRING"
        },
        timeout=(15, 120)
    )

    if response.status_code != 200:
        raise RuntimeError(
            "No se pudo leer el contenido del historico. "
            f"HTTP {response.status_code}. "
            f"Detalle: {response.text[:2000]}"
        )

    return response.json().get(
        "values",
        []
    )


def _find_header_index_in_rows(rows):
    """
    Devuelve el indice 0-based del encabezado real dentro de una
    lista de filas obtenida desde Google Sheets.
    """

    for index, row in enumerate(rows[:25]):
        if _is_expected_header_row(row):
            return index

    return None


def normalize_history_sheet_layout(
    session,
    spreadsheet_id,
    sheet_name=TARGET_SHEET_NAME
):
    """
    Repara historicos creados con la version anterior del codigo.

    Si encuentra metadata antes de los encabezados reales, elimina
    esas filas en Google Sheets para que:

    fila 1 = encabezados
    fila 2+ = registros
    """

    ensure_history_sheet_exists(
        session,
        spreadsheet_id,
        sheet_name
    )

    history_rows = get_history_rows(
        session,
        spreadsheet_id,
        sheet_name
    )

    if not history_rows:
        return history_rows

    header_index = _find_header_index_in_rows(
        history_rows
    )

    if header_index is None:
        return history_rows

    if header_index == 0:
        return history_rows

    metadata_url = (
        "https://sheets.googleapis.com/v4/spreadsheets/"
        f"{spreadsheet_id}"
    )

    metadata_response = session.get(
        metadata_url,
        params={
            "fields": "sheets.properties(sheetId,title)"
        },
        timeout=(15, 60)
    )

    if metadata_response.status_code != 200:
        raise RuntimeError(
            "No se pudo obtener el sheetId para limpiar "
            "el historico. "
            f"HTTP {metadata_response.status_code}. "
            f"Detalle: {metadata_response.text[:2000]}"
        )

    sheet_id = None

    for item in metadata_response.json().get(
        "sheets",
        []
    ):
        properties = item.get(
            "properties",
            {}
        )

        if properties.get("title") == sheet_name:
            sheet_id = properties.get("sheetId")
            break

    if sheet_id is None:
        raise RuntimeError(
            f"No se encontro el sheetId de '{sheet_name}'"
        )

    batch_update_url = (
        "https://sheets.googleapis.com/v4/spreadsheets/"
        f"{spreadsheet_id}:batchUpdate"
    )

    payload = {
        "requests": [
            {
                "deleteDimension": {
                    "range": {
                        "sheetId": sheet_id,
                        "dimension": "ROWS",
                        "startIndex": 0,
                        "endIndex": header_index
                    }
                }
            }
        ]
    }

    response = session.post(
        batch_update_url,
        json=payload,
        timeout=(15, 60)
    )

    if response.status_code != 200:
        raise RuntimeError(
            "No se pudo eliminar la metadata previa "
            "al encabezado del historico. "
            f"HTTP {response.status_code}. "
            f"Detalle: {response.text[:2000]}"
        )

    logger.info(
        "Historico normalizado. "
        "sheet=%s filas_eliminadas=%s",
        sheet_name,
        header_index
    )

    return get_history_rows(
        session,
        spreadsheet_id,
        sheet_name
    )


def _normalize_cell_for_compare(value):
    if value is None:
        return ("empty", "")

    if isinstance(value, bool):
        return ("bool", value)

    if isinstance(value, int):
        return ("number", str(value))

    if isinstance(value, float):
        if value.is_integer():
            return ("number", str(int(value)))

        return (
            "number",
            format(value, ".15g")
        )

    return (
        "text",
        str(value).strip()
    )


def _normalize_row_for_compare(row):
    return tuple(
        _normalize_cell_for_compare(value)
        for value in _trim_row(row)
    )


def append_rows_to_history(
    session,
    spreadsheet_id,
    rows,
    sheet_name=TARGET_SHEET_NAME
):
    """
    Acumula datos en el historico.

    - La primera fila recibida debe ser el encabezado real.
    - Si el historico viejo tiene metadata en A1, la elimina.
    - Se omiten filas que ya existan exactamente en el historico.
    """

    if not rows:
        raise RuntimeError(
            "No hay filas para agregar al historico"
        )

    source_header = rows[0]
    source_data_rows = rows[1:]

    if not _is_expected_header_row(source_header):
        raise RuntimeError(
            "El archivo origen no tiene el encabezado esperado "
            f"para '{sheet_name}'"
        )

    history_rows = normalize_history_sheet_layout(
        session,
        spreadsheet_id,
        sheet_name
    )

    rows_to_append = []

    if not history_rows:
        rows_to_append.append(source_header)
        existing_keys = set()

    else:
        history_header = history_rows[0]

        if (
            _normalize_row_for_compare(history_header)
            !=
            _normalize_row_for_compare(source_header)
        ):
            raise RuntimeError(
                "El encabezado del archivo historico no coincide "
                "con el encabezado actual de "
                f"'{sheet_name}'. "
                "Se detuvo la acumulacion para no mezclar "
                "columnas incorrectamente."
            )

        existing_keys = {
            _normalize_row_for_compare(row)
            for row in history_rows[1:]
            if _trim_row(row)
        }

    added_rows = 0
    skipped_duplicates = 0

    for row in source_data_rows:
        normalized = _normalize_row_for_compare(row)

        if not normalized:
            continue

        if normalized in existing_keys:
            skipped_duplicates += 1
            continue

        rows_to_append.append(row)
        existing_keys.add(normalized)
        added_rows += 1

    if rows_to_append:
        append_url = _sheet_append_url(
            spreadsheet_id,
            sheet_name
        )

        batch_size = 500

        for start in range(
            0,
            len(rows_to_append),
            batch_size
        ):
            batch = rows_to_append[
                start:start + batch_size
            ]

            response = session.post(
                append_url,
                params={
                    "valueInputOption": "RAW",
                    "insertDataOption": "INSERT_ROWS"
                },
                json={
                    "majorDimension": "ROWS",
                    "values": batch
                },
                timeout=(15, 120)
            )

            if response.status_code != 200:
                raise RuntimeError(
                    "Google Sheets API rechazo la escritura "
                    "del historico. "
                    f"HTTP {response.status_code}. "
                    f"Detalle: {response.text[:2000]}"
                )

    previous_data_rows = max(
        len(history_rows) - 1,
        0
    )

    return {
        "source_rows": len(source_data_rows),
        "added_rows": added_rows,
        "skipped_duplicates": skipped_duplicates,
        "historical_rows": (
            previous_data_rows
            + added_rows
        )
    }

def update_history_from_rows(
    rows,
    source_file_id=None
):
    """
    Escribe registros en el Google Sheet configurado en
    GOOGLE_DRIVE_HISTORY_URL.
    """

    if not DEFAULT_HISTORY_URL:
        raise RuntimeError(
            "Falta la variable de entorno "
            "GOOGLE_DRIVE_HISTORY_URL"
        )

    history_file_id = extract_google_file_id(
        DEFAULT_HISTORY_URL
    )

    if (
        source_file_id
        and history_file_id == source_file_id
    ):
        raise RuntimeError(
            "GOOGLE_DRIVE_HISTORY_URL apunta al mismo archivo "
            "que GOOGLE_DRIVE_FILE_URL. El historico debe ser "
            "un Google Sheet diferente."
        )

    credentials = build_google_credentials()
    session = AuthorizedSession(credentials)

    logger.info(
        "Actualizando historico en Google Sheets. "
        "history_file_id=%s service_account=%s",
        history_file_id,
        credentials.service_account_email
    )

    result = append_rows_to_history(
        session,
        history_file_id,
        rows,
        TARGET_SHEET_NAME
    )

    logger.info(
        "Historico actualizado. "
        "source_rows=%s added_rows=%s "
        "duplicates=%s historical_rows=%s",
        result["source_rows"],
        result["added_rows"],
        result["skipped_duplicates"],
        result["historical_rows"]
    )

    return result


# ============================================================
# DESCARGAR HISTORICO COMO XLSX
# ============================================================

def download_history_as_xlsx():
    """
    Descarga el Google Sheet historico como XLSX y devuelve
    solamente la hoja "Comentarios x orden".
    """

    if not DEFAULT_HISTORY_URL:
        raise RuntimeError(
            "Falta la variable de entorno "
            "GOOGLE_DRIVE_HISTORY_URL"
        )

    history_file_id = extract_google_file_id(
        DEFAULT_HISTORY_URL
    )

    credentials = build_google_credentials()
    session = AuthorizedSession(credentials)

    normalize_history_sheet_layout(
        session,
        history_file_id,
        TARGET_SHEET_NAME
    )

    binary_data = download_google_file_as_xlsx_by_id(
        history_file_id,
        session
    )

    filtered_binary = keep_only_target_sheet(
        binary_data,
        TARGET_SHEET_NAME
    )

    return {
        "drive_file_id": history_file_id,
        "data": filtered_binary,
        "size_bytes": len(filtered_binary)
    }


# ============================================================
# PROCESO COMPLETO DEL JOB
# ============================================================

def process_drive_download(job_id):
    """
    Flujo:

    queued
      -> running
      -> descarga origen
      -> conserva solo "Comentarios x orden"
      -> acumula filas nuevas en el historico de Drive
      -> guarda XLSX diario en DB
      -> completed
    """

    job = DriveFileDownload.query.filter_by(
        job_id=job_id
    ).first()

    if not job:
        raise RuntimeError(
            f"No existe el job {job_id}"
        )

    job.status = "running"
    job.started_at = datetime.utcnow()
    job.error = None

    db.session.commit()

    try:
        source_result = download_google_sheet_as_xlsx(
            job.source_url
        )

        source_rows = extract_target_sheet_rows(
            source_result["data"],
            TARGET_SHEET_NAME
        )

        daily_binary = keep_only_target_sheet(
            source_result["data"],
            TARGET_SHEET_NAME
        )

        history_result = update_history_from_rows(
            source_rows,
            source_file_id=source_result["drive_file_id"]
        )

        job.drive_file_id = source_result["drive_file_id"]
        job.data = daily_binary
        job.size_bytes = len(daily_binary)
        job.status = "completed"
        job.finished_at = datetime.utcnow()
        job.error = None

        db.session.commit()

        logger.info(
            "Job Drive completado. "
            "job_id=%s daily_size=%s "
            "history_added=%s history_total=%s",
            job_id,
            len(daily_binary),
            history_result["added_rows"],
            history_result["historical_rows"]
        )

        return {
            "job_id": job_id,
            "status": "completed",
            "size_bytes": len(daily_binary),
            "history": history_result
        }

    except Exception as e:
        db.session.rollback()

        job = DriveFileDownload.query.filter_by(
            job_id=job_id
        ).first()

        if job:
            job.status = "failed"
            job.error = str(e)
            job.finished_at = datetime.utcnow()
            db.session.commit()

        logger.error(
            "Error descargando archivo de Drive. "
            "job_id=%s error=%s",
            job_id,
            str(e),
            exc_info=True
        )

        raise
